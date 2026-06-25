#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
inventario_excel.py
===================
Genera "Viñas Norte - Inventario.xlsx": un inventario de CONTROL (no de
presupuesto). Una hoja por LOTE suministrado más una hoja "Resumen".

Cada hoja de lote tiene, con FÓRMULAS vivas:
  1) EXISTENCIAS: Recibido (lo editas), Salidas (se suma sola de la bitácora)
     y Saldo (= Recibido − Salidas).
  2) BITÁCORA DE SALIDAS: cada salida descuenta sola del saldo (lista de
     material por menú desplegable).
  3) EXTRAS: boquillas, impermeabilizante, etc.
No incluye "requerido"; es puro control de inventario.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# Existencia inicial recibida = lo SUMINISTRADO del presupuesto (fuente única).
# Se puede editar en la propia hoja conforme llega o se ajusta el material.
from datos_cliente import SUMINISTRADO as RECIBIDO, LOTES

MAT = [("Piso Moret Arena",  "caja 1.42 m² (2 pz)",  "Moret"),
       ("Piso Royal Walnut", "caja 1.20 m² (5 pz)",  "Royal"),
       ("Urbania White",     "caja 1.36 m² (10 pz)", "Urbania"),
       ("Malla Lyndhurst",   "pieza 0.30×0.60 m",    "Malla")]
LISTA_MAT = ",".join(m[0] for m in MAT)

EXTRAS = ["Boquilla Cantera", "Boquilla Chocolate", "Boquilla Blanco",
          "Boquilla Plata", "Impermeabilizante para charola", "", ""]

N_BITACORA = 24

HDR = PatternFill("solid", fgColor="1F4E78")
SUB = PatternFill("solid", fgColor="D6E4F0")
EXT = PatternFill("solid", fgColor="FCF3CF")
LOG = PatternFill("solid", fgColor="E8F5E9")
SAL = PatternFill("solid", fgColor="FEF3CF")
WH = Font(color="FFFFFF", bold=True)
B = Font(bold=True)
thin = Side(style="thin", color="BBBBBB")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal="center")


def hoja_lote(ws, lote, mza, modelo):
    rec = RECIBIDO[modelo]
    ws.merge_cells("A1:F1")
    ws["A1"] = f"INVENTARIO Y CONTROL — LOTE {lote}, MANZANA {mza} · {modelo.upper()}"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF"); ws["A1"].fill = HDR
    ws["A1"].alignment = CEN

    # 1) EXISTENCIAS (control con fórmulas)
    ws.cell(3, 1, "1) EXISTENCIAS (CONTROL)").font = B
    enc = ["Material", "Presentación", "Recibido", "Salidas", "Saldo", "Observaciones"]
    for c, h in enumerate(enc, 1):
        cell = ws.cell(4, c, h); cell.font = WH; cell.fill = HDR
        cell.alignment = CEN; cell.border = BORD
    fila_mat0 = 5
    # rango de la bitácora (se define más abajo, pero ya fijamos sus filas)
    bit_hdr = fila_mat0 + len(MAT) + 2          # fila del encabezado de bitácora
    bit0 = bit_hdr + 1
    bit1 = bit0 + N_BITACORA - 1
    for i, (nombre, pres, _key) in enumerate(MAT):
        r = fila_mat0 + i
        ws.cell(r, 1, nombre).font = B
        ws.cell(r, 2, pres)
        # Recibido (editable)
        cr = ws.cell(r, 3, rec[MAT[i][2]]); cr.alignment = CEN
        # Salidas = suma de la bitácora para ese material
        cs = ws.cell(r, 4, f"=SUMIF($B${bit0}:$B${bit1},$A{r},$C${bit0}:$C${bit1})")
        cs.alignment = CEN
        # Saldo = Recibido - Salidas
        cz = ws.cell(r, 5, f"=C{r}-D{r}"); cz.alignment = CEN
        cz.font = B; cz.fill = SAL
        for c in range(1, 7):
            ws.cell(r, c).border = BORD

    # 2) BITÁCORA DE SALIDAS
    ws.cell(bit_hdr - 1, 1, "2) BITÁCORA DE SALIDAS (cada salida descuenta del saldo)").font = B
    benc = ["Fecha", "Material", "Cantidad", "Frente / área", "Recibió en obra"]
    for c, h in enumerate(benc, 1):
        cell = ws.cell(bit_hdr, c, h); cell.font = WH; cell.fill = HDR
        cell.alignment = CEN; cell.border = BORD
    for r in range(bit0, bit1 + 1):
        for c in range(1, 6):
            cell = ws.cell(r, c); cell.border = BORD
            if (r - bit0) % 2 == 1:
                cell.fill = LOG
    # menú desplegable de material en la bitácora (para que el SUMIF cuadre)
    dv = DataValidation(type="list", formula1=f'"{LISTA_MAT}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B{bit0}:B{bit1}")

    # 3) EXTRAS / COMPLEMENTOS
    ex_hdr = bit1 + 2
    ws.cell(ex_hdr, 1, "3) EXTRAS / COMPLEMENTOS").font = B
    for c, h in enumerate(["Material", "Presentación / color", "Cantidad", "Observaciones"], 1):
        cell = ws.cell(ex_hdr + 1, c, h); cell.font = B; cell.fill = SUB; cell.border = BORD
    for i, ex in enumerate(EXTRAS):
        r = ex_hdr + 2 + i
        ws.cell(r, 1, ex)
        for c in range(1, 7):
            ws.cell(r, c).border = BORD; ws.cell(r, c).fill = EXT

    for c, w in enumerate([24, 22, 12, 12, 12, 22], 1):
        ws.column_dimensions[chr(64 + c)].width = w
    # devolver las filas de saldo (para el Resumen)
    return {MAT[i][2]: fila_mat0 + i for i in range(len(MAT))}


def hoja_resumen(ws, refs):
    """Roll-up: lee el Saldo de cada hoja de lote por fórmula."""
    ws.merge_cells("A1:H1")
    ws["A1"] = "RESUMEN DE INVENTARIO — SALDO POR LOTE"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF"); ws["A1"].fill = HDR
    ws["A1"].alignment = CEN
    enc = ["Lote", "Manzana", "Modelo", "Material", "Recibido", "Salidas", "Saldo", "Observaciones"]
    for c, h in enumerate(enc, 1):
        cell = ws.cell(3, c, h); cell.font = WH; cell.fill = HDR
        cell.alignment = CEN; cell.border = BORD
    r = 4
    for (lote, mza, modelo, hname, filas) in refs:
        f0 = r
        for (nombre, _pres, key) in MAT:
            fr = filas[key]
            ws.cell(r, 4, nombre).font = B
            ws.cell(r, 5, f"='{hname}'!C{fr}").alignment = CEN
            ws.cell(r, 6, f"='{hname}'!D{fr}").alignment = CEN
            cz = ws.cell(r, 7, f"='{hname}'!E{fr}"); cz.alignment = CEN; cz.font = B; cz.fill = SAL
            for c in range(1, 9):
                ws.cell(r, c).border = BORD
            r += 1
        for col, val in ((1, f"L{lote}"), (2, f"M{mza}"), (3, modelo)):
            ws.merge_cells(start_row=f0, start_column=col, end_row=r - 1, end_column=col)
            cell = ws.cell(f0, col, val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = B; cell.fill = SUB
        r += 1
    for c, w in enumerate([8, 9, 13, 22, 12, 12, 12, 24], 1):
        ws.column_dimensions[chr(64 + c)].width = w


def main():
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    resumen = wb.create_sheet("Resumen")
    refs = []
    usados = {}
    for (lote, mza, modelo) in LOTES:
        base = f"L{lote} M{mza}"
        nombre = base
        k = 2
        while nombre in usados:           # evita choque de nombres de hoja
            nombre = f"{base} ({k})"; k += 1
        usados[nombre] = True
        ws = wb.create_sheet(nombre)
        filas = hoja_lote(ws, lote, mza, modelo)
        refs.append((lote, mza, modelo, nombre, filas))
    hoja_resumen(resumen, refs)
    out = "Viñas Norte - Inventario.xlsx"
    wb.save(out)
    print("->", out)


if __name__ == "__main__":
    main()
